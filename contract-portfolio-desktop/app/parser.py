# -*- coding: utf-8 -*-
"""
Excel -> Dashboard JSON dönüştürücü.

Kaynak Excel şeması (kolon sırası ve adları):
  Contract Number | Müşteri | Sözleşmenin Yıllık Bedeli | (YTD) SM% |
  (Fiscal Year) SM% | Sözleşme Yenileme(Başlangıç) Tarihi |
  Sözleşme Bitiş Tarihi | Sözleşme Yenileme Period | Sözleşme Bitiş Period |
  (Son Sözleşme Dönemi) SM% | 1. Yıl | 2. Yıl | 3. Yıl | 4. Yıl | 5.yıl |
  Sözleşmenin Toplam Bedeli | Notlar

Çıktı (satır başına, HTML şablonunun beklediği şema):
  {contract, musteri, yillik, tl, ytd, fy, son, start, end, years[5],
   toplam, not, daysToEnd, dataError, expired}

Bu modül HTML'deki hesaplama/grafik mantığına dokunmaz; yalnızca veri katmanını
üretir. daysToEnd / expired hesaplamaları çağrı anındaki güncel tarihe göredir.
"""

from __future__ import annotations

import datetime as _dt
import difflib
import re
from dataclasses import dataclass, field
from typing import Any

try:
    import openpyxl
except ImportError as e:  # pragma: no cover
    raise ImportError(
        "openpyxl kurulu değil. Lütfen 'pip install openpyxl' çalıştırın."
    ) from e


# --------------------------------------------------------------------------- #
# Beklenen kolonlar — sıra ve adlar birebir kaynak dosyadaki gibi.
# --------------------------------------------------------------------------- #
COL_CONTRACT = "Contract Number"
COL_MUSTERI = "Müşteri"
COL_YILLIK = "Sözleşmenin Yıllık Bedeli"
COL_YTD = "(YTD) SM%"
COL_FY = "(Fiscal Year) SM%"
COL_START = "Sözleşme Yenileme(Başlangıç) Tarihi"
COL_END = "Sözleşme Bitiş Tarihi"
COL_RENEW_PERIOD = "Sözleşme Yenileme Period"
COL_END_PERIOD = "Sözleşme Bitiş Period"
COL_SON = "(Son Sözleşme Dönemi) SM%"
COL_Y1 = "1. Yıl"
COL_Y2 = "2. Yıl"
COL_Y3 = "3. Yıl"
COL_Y4 = "4. Yıl"
COL_Y5 = "5.yıl"
COL_TOPLAM = "Sözleşmenin Toplam Bedeli"
COL_NOTLAR = "Notlar"

# Bu kolonlar mutlaka bulunmalı; bulunamazsa kullanıcıya eşleşme önerilir.
REQUIRED_COLUMNS = [
    COL_CONTRACT, COL_MUSTERI, COL_YILLIK, COL_YTD, COL_FY,
    COL_START, COL_END, COL_SON, COL_TOPLAM,
]
# Bu kolonlar opsiyonel (yoksa null / boş kabul edilir).
OPTIONAL_COLUMNS = [
    COL_RENEW_PERIOD, COL_END_PERIOD,
    COL_Y1, COL_Y2, COL_Y3, COL_Y4, COL_Y5, COL_NOTLAR,
]
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS
YEAR_COLUMNS = [COL_Y1, COL_Y2, COL_Y3, COL_Y4, COL_Y5]


class ParseError(Exception):
    """Kullanıcıya gösterilecek, anlaşılır Türkçe mesaj taşıyan hata."""


@dataclass
class ColumnPlan:
    """Kolon eşleştirme planı — hangi beklenen kolon Excel'in kaçıncı sütunu."""
    mapping: dict[str, int] = field(default_factory=dict)      # beklenen ad -> 0-tabanlı sütun index
    exact: list[str] = field(default_factory=list)             # birebir eşleşenler
    suggestions: list[dict] = field(default_factory=list)      # {expected, found, index, score}
    missing_required: list[str] = field(default_factory=list)  # hiç eşleşmeyen zorunlu kolonlar
    missing_optional: list[str] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)           # dosyadaki ham başlıklar

    @property
    def needs_confirmation(self) -> bool:
        return bool(self.suggestions)


@dataclass
class ParseResult:
    rows: list[dict] = field(default_factory=list)
    dq: list[str] = field(default_factory=list)      # veri kalitesi / parse uyarıları (Türkçe)
    plan: ColumnPlan | None = None
    sheet_name: str = ""
    total_rows: int = 0


# --------------------------------------------------------------------------- #
# Normalizasyon / eşleştirme
# --------------------------------------------------------------------------- #
_TR_MAP = str.maketrans({
    "ı": "i", "İ": "i", "ş": "s", "Ş": "s", "ğ": "g", "Ğ": "g",
    "ü": "u", "Ü": "u", "ö": "o", "Ö": "o", "ç": "c", "Ç": "c",
})


def _norm(s: Any) -> str:
    """Başlık karşılaştırması için: küçük harf, TR harf sadeleştirme, boşluk/nokta sadeleştirme."""
    if s is None:
        return ""
    s = str(s).strip().lower().translate(_TR_MAP)
    s = re.sub(r"[\s._\-/()]+", " ", s)  # noktalama ve boşlukları tekilleştir
    return s.strip()


def plan_columns(headers: list[Any]) -> ColumnPlan:
    """Excel başlık satırını beklenen kolonlara eşle. Birebir olmayanlar için öneri üret."""
    plan = ColumnPlan(headers=[("" if h is None else str(h)) for h in headers])
    norm_headers = [_norm(h) for h in headers]
    used_idx: set[int] = set()

    for expected in ALL_COLUMNS:
        ne = _norm(expected)
        idx = None
        # 1) birebir (normalize edilmiş) eşleşme
        for i, nh in enumerate(norm_headers):
            if i in used_idx:
                continue
            if nh == ne:
                idx = i
                break
        if idx is not None:
            plan.mapping[expected] = idx
            plan.exact.append(expected)
            used_idx.add(idx)
            continue

        # 2) bulanık (fuzzy) en yakın eşleşme
        candidates = [(i, nh) for i, nh in enumerate(norm_headers) if i not in used_idx and nh]
        best_i, best_score = None, 0.0
        for i, nh in candidates:
            score = difflib.SequenceMatcher(None, ne, nh).ratio()
            # bir taraf diğerini içeriyorsa puanı yükselt (ör. "musteri" / "musteri adi")
            if ne in nh or nh in ne:
                score = max(score, 0.9)
            if score > best_score:
                best_i, best_score = i, score
        if best_i is not None and best_score >= 0.62:
            plan.mapping[expected] = best_i
            plan.suggestions.append({
                "expected": expected,
                "found": plan.headers[best_i],
                "index": best_i,
                "score": round(best_score, 2),
            })
            used_idx.add(best_i)
        else:
            if expected in REQUIRED_COLUMNS:
                plan.missing_required.append(expected)
            else:
                plan.missing_optional.append(expected)

    return plan


# --------------------------------------------------------------------------- #
# Hücre değeri ayrıştırıcılar
# --------------------------------------------------------------------------- #
_CURRENCY_TL_HINTS = ("₺", "try", "tl")
_CURRENCY_EUR_HINTS = ("€", "eur")


def parse_money(value: Any) -> tuple[float | None, bool]:
    """
    Para değeri ayrıştır. '€ 40,000.00', '₺1,200,000.00', 40000, '40.000,00' vb.
    Dönüş: (sayı_ya_da_None, tl_mi)
    """
    if value is None:
        return None, False
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value), False

    s = str(value).strip()
    if not s:
        return None, False

    low = s.lower()
    is_tl = any(h in low for h in _CURRENCY_TL_HINTS)
    # '€'/'eur' varsa kesin EUR; sembol yoksa varsayılan EUR (tl=False)
    if any(h in low for h in _CURRENCY_EUR_HINTS):
        is_tl = False

    # Sayı gövdesini çıkar: rakam, virgül, nokta, eksi
    body = re.sub(r"[^\d,.\-]", "", s)
    if body in ("", "-", ".", ","):
        return None, is_tl

    num = _parse_number_body(body)
    return num, is_tl


def _parse_number_body(body: str) -> float | None:
    """'40,000.00' / '40.000,00' / '40000' / '21.04' gibi gövdeleri float'a çevir."""
    neg = body.startswith("-")
    body = body.lstrip("-")
    has_comma = "," in body
    has_dot = "." in body

    if has_comma and has_dot:
        # Son görünen ayraç ondalık ayraçtır.
        if body.rfind(",") > body.rfind("."):
            # virgül ondalık (Avrupa): noktaları at, virgülü noktaya çevir
            body = body.replace(".", "").replace(",", ".")
        else:
            # nokta ondalık (US/fatura): virgülleri (binlik) at
            body = body.replace(",", "")
    elif has_comma:
        # Yalnızca virgül: 1-2 basamak sonrası ondalık, değilse binlik
        frac = body.split(",")[-1]
        if len(body.split(",")) == 2 and 1 <= len(frac) <= 2:
            body = body.replace(",", ".")
        else:
            body = body.replace(",", "")
    elif has_dot:
        # Yalnızca nokta: birden çok nokta => binlik ayraç
        if body.count(".") > 1:
            body = body.replace(".", "")
        # tek nokta => ondalık, dokunma
    try:
        val = float(body)
    except ValueError:
        return None
    return -val if neg else val


def parse_percent(value: Any) -> tuple[float | None, str | None]:
    """
    SM% değerini ayrıştır. '% 21.04%', 21.04, 'x/o', boş vb.
    Dönüş: (sayı_ya_da_None, hata_metni_ya_da_None)
    Not: 0-100 ölçeği korunur; yalnızca hücre yüzde-biçimli float ise *100 yapılır
         (bu ayrım openpyxl number_format ile parse_row içinde ele alınır).
    """
    if value is None:
        return None, None
    if isinstance(value, bool):
        return None, "mantıksal (TRUE/FALSE) değer"
    if isinstance(value, (int, float)):
        return float(value), None

    s = str(value).strip()
    if s == "":
        return None, None
    # yüzde işareti / boşlukları temizle
    cleaned = s.replace("%", "").replace(" ", "").replace(" ", "")
    cleaned = cleaned.replace(",", ".")  # olası ondalık virgül
    m = re.fullmatch(r"-?\d+(\.\d+)?", cleaned)
    if not m:
        # 'x/o', 'x', '-', 'n/a' gibi metinsel değerler
        return None, f'metin değer ("{s}")'
    try:
        return float(cleaned), None
    except ValueError:
        return None, f'okunamayan değer ("{s}")'


_DATE_TEXT_FORMATS = (
    "%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y",
    "%d.%m.%y", "%d/%m/%y", "%m/%d/%Y",
)


def parse_date(value: Any) -> tuple[_dt.date | None, str | None]:
    """
    Tarih ayrıştır. '01.01.2024', '10/05/2022', ISO, ya da Excel datetime.
    Dönüş: (date_ya_da_None, hata_metni_ya_da_None)
    """
    if value is None:
        return None, None
    if isinstance(value, _dt.datetime):
        return value.date(), None
    if isinstance(value, _dt.date):
        return value, None

    s = str(value).strip()
    if s == "":
        return None, None
    # Sadece tarih kısmını al (saat varsa at)
    s = s.split(" ")[0].split("T")[0]
    for fmt in _DATE_TEXT_FORMATS:
        try:
            return _dt.datetime.strptime(s, fmt).date(), None
        except ValueError:
            continue
    return None, f'okunamayan tarih ("{value}")'


def _clean_contract(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


# --------------------------------------------------------------------------- #
# Ana ayrıştırma
# --------------------------------------------------------------------------- #
def parse_workbook(
    path: str,
    plan: ColumnPlan | None = None,
    today: _dt.date | None = None,
    sheet_name: str | None = None,
) -> ParseResult:
    """
    Excel dosyasını oku ve dashboard satırlarını üret.

    plan verilmezse başlık satırından otomatik plan çıkarılır. Zorunlu kolon
    eksikse ParseError yükseltilir. Bulanık eşleşme önerileri varsa (plan
    verilmediyse) yine de devam edilir; onay akışını çağıran katman yönetir.
    """
    if today is None:
        today = _dt.date.today()

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        raise ParseError(
            f"Excel dosyası açılamadı: {e}\n"
            "Dosyanın .xlsx biçiminde ve bozuk olmadığından emin olun."
        ) from e

    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    if ws is None:
        raise ParseError("Çalışma kitabında okunabilir bir sayfa bulunamadı.")

    rows_iter = ws.iter_rows(values_only=False)
    try:
        header_cells = next(rows_iter)
    except StopIteration:
        raise ParseError("Excel dosyası boş görünüyor (başlık satırı yok).")

    headers = [c.value for c in header_cells]
    if plan is None:
        plan = plan_columns(headers)

    if plan.missing_required:
        eksik = ", ".join(plan.missing_required)
        bulunan = ", ".join(h for h in plan.headers if h) or "(başlık yok)"
        raise ParseError(
            "Zorunlu kolon(lar) bulunamadı:\n"
            f"  • {eksik}\n\n"
            f"Dosyadaki başlıklar:\n  {bulunan}\n\n"
            "Lütfen kolon adlarını beklenen şemaya göre düzeltip tekrar deneyin."
        )

    result = ParseResult(plan=plan, sheet_name=ws.title)

    # number_format'a bakabilmek için hücre nesnelerini kullanıyoruz.
    def cell_of(row_cells, expected_col):
        idx = plan.mapping.get(expected_col)
        if idx is None or idx >= len(row_cells):
            return None
        return row_cells[idx]

    def val_of(row_cells, expected_col):
        c = cell_of(row_cells, expected_col)
        return None if c is None else c.value

    excel_row_no = 1  # başlık satırı = 1
    for row_cells in rows_iter:
        excel_row_no += 1
        # Tamamen boş satırları atla
        if all((c.value is None or str(c.value).strip() == "") for c in row_cells):
            continue

        try:
            row = _parse_row(row_cells, plan, cell_of, val_of, excel_row_no, today, result.dq)
        except Exception as e:
            # Hiçbir satır uygulamayı düşürmemeli.
            result.dq.append(
                f"Satır {excel_row_no}: beklenmeyen bir hata nedeniyle atlandı ({e})."
            )
            continue
        if row is not None:
            result.rows.append(row)

    result.total_rows = len(result.rows)
    wb.close()

    if result.total_rows == 0:
        raise ParseError(
            "Dosyada veri satırı bulunamadı. Başlık satırının altında en az bir "
            "sözleşme satırı olmalı."
        )
    return result


def _pct_from_cell(cell, dq: list[str], row_no: int, col_name: str) -> float | None:
    """Bir SM% hücresini 0-100 ölçeğinde sayıya çevir; metinleri null yapıp DQ'ya ekle."""
    if cell is None:
        return None
    num, err = parse_percent(cell.value)
    if err is not None:
        dq.append(f'Satır {row_no}, "{col_name}": {err} → boş kabul edildi.')
        return None
    if num is None:
        return None
    # Hücre yüzde biçimliyse (0.21 gibi) 0-100 ölçeğine taşı.
    fmt = getattr(cell, "number_format", "") or ""
    if "%" in fmt and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
        num *= 100.0
    return round(num, 2)


def _parse_row(row_cells, plan, cell_of, val_of, row_no, today, dq) -> dict | None:
    contract = _clean_contract(val_of(row_cells, COL_CONTRACT))
    musteri = val_of(row_cells, COL_MUSTERI)
    musteri = "" if musteri is None else str(musteri).strip()

    # Para: yıllık + toplam (TL tespiti bunlardan)
    yillik, tl_y = parse_money(val_of(row_cells, COL_YILLIK))
    toplam, tl_t = parse_money(val_of(row_cells, COL_TOPLAM))
    tl = bool(tl_y or tl_t)

    if yillik is None:
        dq.append(f'Satır {row_no}, "{COL_YILLIK}": sayıya çevrilemedi → 0 kabul edildi.')
        yillik = 0.0

    # SM% kolonları
    ytd = _pct_from_cell(cell_of(row_cells, COL_YTD), dq, row_no, COL_YTD)
    fy = _pct_from_cell(cell_of(row_cells, COL_FY), dq, row_no, COL_FY)
    son = _pct_from_cell(cell_of(row_cells, COL_SON), dq, row_no, COL_SON)

    # Tarihler
    start, err_s = parse_date(val_of(row_cells, COL_START))
    end, err_e = parse_date(val_of(row_cells, COL_END))
    if err_s:
        dq.append(f'Satır {row_no}, "{COL_START}": {err_s} → boş kabul edildi.')
    if err_e:
        dq.append(f'Satır {row_no}, "{COL_END}": {err_e} → boş kabul edildi.')

    # Yıl kolonları (çoğunlukla boş)
    years: list[float | None] = []
    for yc in YEAR_COLUMNS:
        if yc in plan.mapping:
            v, _tl = parse_money(val_of(row_cells, yc))
            years.append(v)
        else:
            years.append(None)
    years = (years + [None] * 5)[:5]

    notlar = val_of(row_cells, COL_NOTLAR)
    notlar = "" if notlar is None else str(notlar).strip()

    # Türetilen alanlar — HTML'deki parse.py mantığıyla aynı, güncel tarihe göre.
    data_error = bool(start and end and end < start)
    days_to_end = (end - today).days if end else None
    expired = bool(end and end < today and not data_error)

    return {
        "contract": contract,
        "musteri": musteri,
        "yillik": yillik,
        "tl": tl,
        "ytd": ytd,
        "fy": fy,
        "son": son,
        "start": start.isoformat() if start else None,
        "end": end.isoformat() if end else None,
        "years": years,
        "toplam": toplam if toplam is not None else 0.0,
        "not": notlar,
        "daysToEnd": days_to_end,
        "dataError": data_error,
        "expired": expired,
    }
